#!/usr/bin/env python3
"""Fabrique les documents vierges du menu « Créer nouveau » de Dolphin (clic droit → Créer nouveau).

Trois fichiers Microsoft Office vides, au format de la PME française (A4, Calibri 11 : Carlito les rend à
l'identique, voir docs/compatibilite.md), placés dans system_files/usr/share/templates/.source/ avec leur
lanceur .desktop. Les fichiers sont versionnés : ce script ne sert qu'à les refaire.

Usage : branding/fabriquer-modeles.py
Prérequis (hors BinixX OS, dans un environnement temporaire) : pip install python-docx openpyxl python-pptx
"""

import datetime
import pathlib

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.oxml.ns import qn
from docx.shared import Cm, Pt
from openpyxl import Workbook
from pptx import Presentation
from pptx.util import Inches

RACINE = pathlib.Path(__file__).resolve().parents[1]
DOSSIER = RACINE / "system_files/usr/share/templates"
SOURCE = DOSSIER / ".source"
DATE = datetime.datetime(2026, 1, 1)

LANCEURS = {
    "Document-texte.desktop": ("Document texte", "Document.docx", "x-office-document"),
    "Classeur.desktop": ("Classeur", "Classeur.xlsx", "x-office-spreadsheet"),
    "Presentation.desktop": ("Présentation", "Presentation.pptx", "x-office-presentation"),
}


def document():
    doc = Document()
    section = doc.sections[0]
    section.orientation = WD_ORIENT.PORTRAIT
    section.page_width, section.page_height = Cm(21.0), Cm(29.7)
    section.left_margin = section.right_margin = Cm(2.5)
    section.top_margin = section.bottom_margin = Cm(2.5)
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    # Police aussi pour l'Asie et les polices complexes, et langue du texte : français
    rpr = normal.element.get_or_add_rPr()
    fonts = rpr.get_or_add_rFonts()
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        fonts.set(qn(attr), "Calibri")
    lang = rpr.makeelement(qn("w:lang"), {qn("w:val"): "fr-FR"})
    rpr.append(lang)
    doc.core_properties.author = ""
    doc.core_properties.created = doc.core_properties.modified = DATE
    doc.core_properties.title = ""
    doc.save(SOURCE / "Document.docx")


def classeur():
    wb = Workbook()
    wb.active.title = "Feuil1"
    wb.properties.creator = ""
    wb.properties.created = wb.properties.modified = DATE
    wb.save(SOURCE / "Classeur.xlsx")


def presentation():
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)  # 16:9, comme PowerPoint aujourd'hui
    prs.core_properties.author = ""
    prs.core_properties.created = prs.core_properties.modified = DATE
    prs.core_properties.title = ""
    prs.save(SOURCE / "Presentation.pptx")


def lanceurs():
    for nom, (titre, fichier, icone) in LANCEURS.items():
        (DOSSIER / nom).write_text(
            f"[Desktop Entry]\nName={titre}\nType=Link\nURL=.source/{fichier}\nIcon={icone}\n", encoding="utf-8")


if __name__ == "__main__":
    SOURCE.mkdir(parents=True, exist_ok=True)
    document()
    classeur()
    presentation()
    lanceurs()
    for chemin in sorted(DOSSIER.rglob("*")):
        if chemin.is_file():
            print(f"{chemin.relative_to(RACINE)} ({chemin.stat().st_size} octets)")
